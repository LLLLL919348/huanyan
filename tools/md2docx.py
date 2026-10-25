"""把 docs/安全评估报告-幻言.md 转成可直接打印的 docx。
只用 python-docx；不引入 pandoc，规则很窄但够用：
  # / ## / ### 标题、| 表格 |、- 列表、> 引用、普通段落、--- 分隔线。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

MD = Path(__file__).resolve().parent.parent / 'docs' / '安全评估报告-幻言.md'
OUT = MD.with_suffix('.docx')

doc = Document()

# 基础样式：中文用宋体，正文 11pt，1.5 倍行距
normal = doc.styles['Normal']
normal.font.name = '宋体'
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(4)

for sec in doc.sections:
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(3.0)


def add_heading(text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    sizes = {1: 18, 2: 14, 3: 12}
    run.font.size = Pt(sizes.get(level, 12))
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
    p.paragraph_format.space_after = Pt(6)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_rich(par, text):
    """处理 **粗体**、`代码`、【】高亮"""
    text = text.replace('`', '')
    for part in re.split(r'(\*\*[^*]+\*\*)', text):
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        seg = part[2:-2] if bold else part
        for sub in re.split(r'(【[^】]*】)', seg):
            if not sub:
                continue
            run = par.add_run(sub)
            run.bold = bold
            if sub.startswith('【') and sub.endswith('】'):
                run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
                run.bold = True


def add_table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, cell in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        par = c.paragraphs[0]
        run = par.add_run(cell)
        run.bold = True
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in body:
        cells = t.add_row().cells
        for i, val in enumerate(row[:len(header)]):
            cells[i].text = ''
            add_rich(cells[i].paragraphs[0], val)
    doc.add_paragraph()


lines = MD.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    stripped = line.strip()

    if not stripped:
        i += 1
        continue

    if stripped == '---':
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run('— — —').font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        i += 1
        continue

    m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
    if m:
        add_heading(m.group(2), len(m.group(1)))
        i += 1
        continue

    if stripped.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            raw = lines[i].strip().strip('|')
            if re.match(r'^[\s:|-]+$', raw):
                i += 1
                continue
            rows.append([c.strip() for c in raw.split('|')])
            i += 1
        if rows:
            add_table(rows)
        continue

    if stripped.startswith('>'):
        par = doc.add_paragraph()
        par.paragraph_format.left_indent = Cm(0.8)
        add_rich(par, stripped.lstrip('> ').strip())
        for run in par.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        i += 1
        continue

    m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
    if m:
        digits = m.group(2).rstrip('.')
        style = 'List Number' if digits.isdigit() else 'List Bullet'
        par = doc.add_paragraph(style=style)
        add_rich(par, m.group(3))
        i += 1
        continue

    par = doc.add_paragraph()
    add_rich(par, stripped)
    i += 1

doc.save(OUT)
print('已生成:', OUT)
print('段落数:', len(doc.paragraphs), '表格数:', len(doc.tables))
